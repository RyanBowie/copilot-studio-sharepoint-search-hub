import base64
import binascii
import hashlib
import io
import json
from pathlib import Path
import re
import struct
import sys
import unittest
import xml.etree.ElementTree as ET
import zipfile

import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "src"
AGENT = ROOT.parent / "agent"
SOLUTION_ARCHIVE = ROOT / "CorpNetSearchHubReference_1_0_0_0_unmanaged.zip"
FLOW_ID = "96794fbd-20ae-f111-aaab-002248403bef"
sys.path.insert(0, str(ROOT / "tools"))
from sync_reference import compiled_component


def walk(value):
    yield value
    if isinstance(value, dict):
        for child in value.values():
            yield from walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from walk(child)


def xml_structure(node):
    return (node.tag, sorted(node.attrib.items()), (node.text or "").strip(),
            [xml_structure(child) for child in node])


class SolutionReferenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.flow_path, = (SOURCE / "Workflows").glob("*.json")
        cls.flow = json.loads(cls.flow_path.read_text(encoding="utf-8"))
        cls.definition = cls.flow["properties"]["definition"]
        cls.components = {
            p.parent.name: yaml.safe_load(p.read_text(encoding="utf-8"))
            for p in (SOURCE / "botcomponents").glob("*/data")
        }
        cls.customizations = ET.parse(SOURCE / "Other" / "Customizations.xml")

    def test_all_source_documents_parse(self):
        for path in SOURCE.rglob("*"):
            if path.suffix == ".xml":
                ET.parse(path)
            elif path.suffix == ".json":
                json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(len(self.components), 16)
        for data in self.components.values():
            self.assertNotIn("mcs.metadata", data)
            self.assertIn(data["kind"], ("AdaptiveDialog", "TaskDialog", "GptComponentMetadata"))

    def test_portable_components_match_after_explicit_identity_binding(self):
        for schema, component in self.components.items():
            self.assertEqual(component, compiled_component(schema), schema)
        self.assertEqual(
            self.components["cnh_corpnetSearchHub.gpt.default"]["aISettings"]["model"]["modelNameHint"],
            "GPT5Chat")

    def test_native_definition_is_exact_standing_portable_source(self):
        portable = json.loads((AGENT / "flows" / "search-export" / "definition.json").read_text(encoding="utf-8"))
        self.assertEqual(self.definition, portable)

    def test_both_packaged_search_requests_use_document_id_only_order(self):
        requests = []
        for node in walk(self.definition):
            if isinstance(node, dict):
                for name in ("Initial_request", "Next_request"):
                    action = node.get(name)
                    if isinstance(action, dict) and action.get("type") == "Compose":
                        requests.append(action["inputs"]["request"])
        self.assertEqual(len(requests), 2)
        for request in requests:
            self.assertEqual(request["SortList"], [{"Property": "[docid]", "Direction": 0}])
            self.assertEqual(request["RowLimit"], 100)

    def test_every_flow_reference_resolves_to_the_single_packaged_flow(self):
        refs = [node["flowId"] for component in self.components.values()
                for node in walk(component) if isinstance(node, dict) and "flowId" in node]
        self.assertEqual(refs, [FLOW_ID, FLOW_ID])
        manifest = ET.parse(SOURCE / "Other" / "Solution.xml")
        roots = manifest.findall(".//RootComponent")
        self.assertEqual([(v.attrib["type"], v.attrib["id"].strip("{}").lower()) for v in roots],
                         [("29", FLOW_ID)])
        association = ET.parse(SOURCE / "Assets" / "botcomponent_workflowset.xml").getroot()[0]
        self.assertEqual(association.attrib["workflowid.workflowid"], FLOW_ID)
        self.assertIn(association.attrib["botcomponentid.schemaname"], self.components)
        self.assertEqual(len(list((SOURCE / "Workflows").glob("*.json"))), 1)

    def test_bot_parent_and_active_component_membership(self):
        for path in (SOURCE / "botcomponents").glob("*/botcomponent.xml"):
            component = ET.parse(path).getroot()
            self.assertIn(component.attrib["schemaname"], self.components)
            self.assertEqual(component.findtext("parentbotid/schemaname"), "cnh_corpnetSearchHub")
            self.assertEqual(component.findtext("statecode"), "0")
        self.assertEqual(len(list((SOURCE / "Assets").iterdir())), 1)
        self.assertFalse(any("Untitled" in name or "SendanHTTP" in name for name in self.components))

    def test_native_entity_fragments_are_element_first_and_match_source_bytes(self):
        fragments = list((SOURCE / "bots").glob("*/bot.xml"))
        fragments += list((SOURCE / "botcomponents").glob("*/botcomponent.xml"))
        self.assertEqual(len(fragments), 17)
        with zipfile.ZipFile(SOLUTION_ARCHIVE) as archive:
            for path in fragments:
                entry = path.relative_to(SOURCE).as_posix()
                with self.subTest(entry=entry):
                    content = archive.read(entry)
                    # Native source-control entity files are fragments, not XML documents.
                    self.assertRegex(content, rb"^\s*<" + path.stem.encode("ascii") + rb"(?:\s|>)")
                    self.assertEqual(content, path.read_bytes())

    def test_five_invoker_references_have_no_connected_account_ids(self):
        references = self.flow["properties"]["connectionReferences"]
        self.assertEqual(set(references), {
            "shared_sharepointonline", "shared_office365", "shared_office365users",
            "shared_onedriveforbusiness", "shared_excelonlinebusiness"})
        declared = {r.attrib["connectionreferencelogicalname"]
                    for r in self.customizations.findall(".//connectionreference")}
        used = set()
        for reference in references.values():
            self.assertEqual(reference["runtimeSource"], "invoker")
            self.assertEqual(set(reference["connection"]), {"connectionReferenceLogicalName"})
            used.add(reference["connection"]["connectionReferenceLogicalName"])
        self.assertEqual(used, declared)
        self.assertEqual(len(declared), 5)

    def test_deployment_settings_are_unbound_and_have_no_current_environment_values(self):
        settings = json.loads((ROOT / "deployment-settings.template.json").read_text(encoding="utf-8-sig"))
        self.assertEqual(settings["EnvironmentVariables"], [])
        self.assertEqual(len(settings["ConnectionReferences"]), 5)
        self.assertTrue(all(v["ConnectionId"] == "" for v in settings["ConnectionReferences"]))
        self.assertEqual(settings["CopilotAgents"], [{
            "AadGroupId": "00000000-0000-0000-0000-000000000000", "Name": "cnh_corpnetSearchHub"}])

    def test_import_does_not_publish_enable_channels_or_activate_the_flow(self):
        config = json.loads((SOURCE / "bots" / "cnh_corpnetSearchHub" / "configuration.json").read_text())
        self.assertIs(config["publishOnImport"], False)
        self.assertEqual(config["channels"], [])
        self.assertIs(config["settings"]["GenerativeActionsEnabled"], False)
        metadata = ET.parse(next((SOURCE / "Workflows").glob("*.data.xml")))
        self.assertEqual(metadata.findtext("StateCode"), "0")
        self.assertEqual(metadata.findtext("StatusCode"), "1")
        bot = ET.parse(SOURCE / "bots" / "cnh_corpnetSearchHub" / "bot.xml")
        self.assertEqual(bot.findtext("authenticationmode"), "2")
        self.assertEqual(bot.findtext("authenticationtrigger"), "1")
        self.assertIsNone(bot.find("iconbase64"))

    def test_only_reviewed_intrinsic_or_fictional_guids_occur(self):
        allowed = {FLOW_ID} | {f"00000000-0000-4000-8000-{i:012d}" for i in range(1, 5)}
        pattern = r"(?i)\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b"
        for path in SOURCE.rglob("*"):
            if path.is_file():
                matches = re.findall(pattern, str(path.relative_to(SOURCE)) + path.read_text(encoding="utf-8-sig"))
                self.assertLessEqual({v.lower() for v in matches}, allowed, str(path))

    def test_no_connected_tenant_locators_or_secret_url_shapes(self):
        for path in SOURCE.rglob("*"):
            if not path.is_file():
                continue
            text = path.read_text(encoding="utf-8-sig")
            self.assertNotRegex(text.lower(), r"mngenv|org[0-9a-f]+\.crm\d*\.dynamics\.com|corpnetsearchpoc-")
            self.assertNotRegex(text, r"(?i)[?&](sig|access_token|client_secret)=")
            self.assertNotRegex(text, r"(?i)https://[^/'\"\s]+\.onmicrosoft\.com")
            hosts = re.findall(r"https://([A-Za-z0-9.-]+)(?=/|['\"])", text)
            self.assertLessEqual(set(hosts), {
                "adaptivecards.io", "contoso.sharepoint.com", "contoso-my.sharepoint.com",
                "schema.management.azure.com", "www.w3.org"})

    def test_embedded_workbook_is_the_exact_blank_portable_template(self):
        binaries = [base64.b64decode(value["$content"]) for value in walk(self.definition)
                    if isinstance(value, dict) and "$content" in value]
        self.assertEqual(len(binaries), 1)
        expected = (AGENT / "flows" / "search-export" / "RuntimeSearchResults.template.xlsx").read_bytes()
        self.assertEqual(binaries[0], expected)
        with zipfile.ZipFile(io.BytesIO(binaries[0])) as workbook:
            self.assertIsNone(workbook.testzip())

    def test_actual_solution_zip_is_unmanaged_and_contains_the_current_source(self):
        with zipfile.ZipFile(SOLUTION_ARCHIVE) as archive:
            self.assertIsNone(archive.testzip())
            self.assertEqual(len(archive.namelist()), len(set(archive.namelist())))
            manifest = ET.fromstring(archive.read("solution.xml"))
            self.assertEqual(manifest.findtext(".//Managed"), "0")
            self.assertEqual(manifest.findtext(".//UniqueName"), "cnh_CorpNetSearchHubReference")
            self.assertEqual(len(manifest.findall(".//MissingDependencies/*")), 0)
            for path in SOURCE.rglob("*"):
                if not path.is_file() or path.suffix == ".xml":
                    continue
                entry = path.relative_to(SOURCE).as_posix()
                self.assertEqual(archive.read(entry), path.read_bytes(), entry)
            self.assertNotIn("Assets/botcomponent_connectionreferenceset.xml", archive.namelist())

    def test_pac_generated_xml_preserves_manifest_and_workflow_source_metadata(self):
        with zipfile.ZipFile(SOLUTION_ARCHIVE) as archive:
            self.assertEqual(xml_structure(ET.fromstring(archive.read("solution.xml"))),
                             xml_structure(ET.parse(SOURCE / "Other" / "Solution.xml").getroot()))
            customizations = ET.fromstring(archive.read("customizations.xml"))
            workflows = customizations.findall(".//Workflow")
            self.assertEqual(len(workflows), 1)
            self.assertEqual(xml_structure(workflows[0]),
                             xml_structure(ET.parse(next((SOURCE / "Workflows").glob("*.data.xml"))).getroot()))
            for parent in customizations.iter():
                if workflows[0] in list(parent):
                    parent.remove(workflows[0])
                    break
            self.assertEqual(xml_structure(customizations), xml_structure(self.customizations.getroot()))

    def test_setup_guide_covers_every_compiled_fictional_tenant_location(self):
        locations = {
            "Approved_scopes": ("inputs",),
            "SharePoint_profile": ("inputs", "parameters", "dataset"),
            "Initial_search": ("inputs", "parameters", "dataset"),
            "Next_search_page": ("inputs", "parameters", "dataset"),
            "Personal_site_available": ("expression",),
            "Remember_report_url": ("inputs", "value"),
            "Preview_basic_locators": ("inputs", "where"),
            "Preview_item_verified": ("expression",),
            "Preview_row": ("inputs", "URL"),
            "Basic_locators": ("inputs", "where"),
            "Select_verified_rows": ("inputs", "select", "URL"),
        }
        guide = (ROOT / "README.md").read_text(encoding="utf-8")
        for action, suffix in locations.items():
            self.assertIn("`" + ".".join((action, *suffix)) + "`", guide)
        found = set()

        def inspect(value, path=()):
            if isinstance(value, dict):
                for key, child in value.items():
                    inspect(child, (*path, key))
            elif isinstance(value, list):
                for index, child in enumerate(value):
                    inspect(child, (*path, str(index)))
            elif isinstance(value, str) and re.search(
                    r"contoso|00000000-0000-4000-8000-00000000000[1-4]", "/".join(path) + value):
                matches = [name for name, suffix in locations.items()
                           if any(path[index:index + len(suffix) + 2] == ("actions", name, *suffix)
                                  for index in range(len(path)))]
                self.assertEqual(len(matches), 1, path)
                found.add(matches[0])

        inspect(self.definition)
        self.assertEqual(found, set(locations))

    def test_decoded_workbook_has_only_blank_template_content_and_internal_relationships(self):
        payload, = [base64.b64decode(node["$content"], validate=True) for node in walk(self.definition)
                    if isinstance(node, dict) and "$content" in node]
        with zipfile.ZipFile(io.BytesIO(payload)) as workbook:
            self.assertEqual(set(workbook.namelist()), {
                "docProps/app.xml", "docProps/core.xml", "xl/theme/theme1.xml",
                "xl/worksheets/sheet1.xml", "xl/comments/comment1.xml",
                "xl/drawings/commentsDrawing1.vml", "xl/tables/table1.xml",
                "xl/worksheets/_rels/sheet1.xml.rels", "xl/worksheets/sheet2.xml",
                "xl/tables/table2.xml", "xl/worksheets/_rels/sheet2.xml.rels",
                "xl/styles.xml", "_rels/.rels", "xl/workbook.xml",
                "xl/_rels/workbook.xml.rels", "[Content_Types].xml",
            })
            self.assertEqual(len(workbook.namelist()), len(set(workbook.namelist())))
            for name in workbook.namelist():
                self.assertTrue(name.endswith((".xml", ".rels", ".vml")), name)
                self.assertNotRegex(name, r"(?i)externalLinks|connections|queryTables|embeddings|vbaProject")
                text = workbook.read(name).decode("utf-8")
                self.assertNotRegex(text, r"(?i)\.sharepoint\.com|\.onmicrosoft\.com|[A-Z]:\\Users\\|access_token|client_secret")
                root = ET.fromstring(text)
                if name.endswith(".rels"):
                    self.assertTrue(all(node.get("TargetMode") != "External" for node in root))
                if name == "docProps/core.xml":
                    ns = {"dc": "http://purl.org/dc/elements/1.1/"}
                    self.assertEqual(root.findtext("dc:creator", namespaces=ns), "CorpNet Search Hub")
                if name.startswith("xl/comments/"):
                    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    self.assertEqual([node.text for node in root.findall("s:authors/s:author", ns)],
                                     ["CorpNet Search Hub"])
        book = load_workbook(io.BytesIO(payload))
        try:
            self.assertEqual(book.sheetnames, ["Results", "ExportInfo"])
            results = book["Results"]
            self.assertEqual(results.tables["SearchResults"].ref, "A8:K9")
            self.assertEqual(results.max_row, 9)
            self.assertEqual(results["A9"].value, "__CORPNET_EMPTY_EXPORT__")
            self.assertEqual({cell.coordinate for row in results for cell in row if cell.data_type == "f"},
                             {"B4", "B5", "G9", "I9", "J9"})
            self.assertTrue(all(results[cell].value is None for cell in ("B9", "C9", "D9", "E9", "F9", "H9", "K9")))
            info = book["ExportInfo"]
            self.assertEqual(info.tables["ExportMetadata"].ref, "A1:B24")
            self.assertEqual(info.max_row, 24)
            self.assertEqual(info["A6"].value, "CompletionStatus")
            self.assertEqual(info["B6"].value, "Not started")
            self.assertTrue(all(info.cell(row, 2).value is None for row in range(2, 25) if row != 6))
            self.assertTrue(all(cell.hyperlink is None for sheet in book for row in sheet for cell in row))
            self.assertFalse(book._external_links)
        finally:
            book.close()

    def test_inline_pngs_match_reviewed_assets_without_hidden_metadata(self):
        manifest = json.loads((AGENT / "cards" / "assets.json").read_text())
        assets = list(manifest["themes"].values()) + [manifest["thirdParty"]["assets/sharepoint-48.png"]]
        expected = {asset["sha256"] for asset in assets}
        encoded = set(re.findall(r"data:image/png;base64,([A-Za-z0-9+/=]+)", json.dumps(self.components)))
        self.assertEqual(len(encoded), 6)
        found = set()
        for value in encoded:
            payload = base64.b64decode(value, validate=True)
            found.add(hashlib.sha256(payload).hexdigest())
            self.assertEqual(payload[:8], b"\x89PNG\r\n\x1a\n")
            position, chunks = 8, []
            while position < len(payload):
                length, = struct.unpack(">I", payload[position:position + 4])
                kind = payload[position + 4:position + 8]
                data_end = position + 8 + length
                self.assertLessEqual(data_end + 4, len(payload))
                crc, = struct.unpack(">I", payload[data_end:data_end + 4])
                self.assertEqual(crc, binascii.crc32(payload[position + 4:data_end]) & 0xffffffff)
                chunks.append(kind)
                position = data_end + 4
                if kind == b"IEND":
                    break
            self.assertEqual(position, len(payload), "Unexpected trailing PNG data")
            self.assertEqual(chunks[0], b"IHDR")
            self.assertEqual(chunks[-1], b"IEND")
            self.assertLessEqual(set(chunks), {b"IHDR", b"PLTE", b"tRNS", b"IDAT", b"IEND", b"pHYs"})
        self.assertEqual(found, expected)


if __name__ == "__main__":
    unittest.main()
