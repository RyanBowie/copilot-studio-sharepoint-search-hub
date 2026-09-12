import copy
import importlib.util
import io
import json
from pathlib import Path
import re
import sys
import shutil
import unittest
from unittest.mock import patch
import uuid

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
SPEC = importlib.util.spec_from_file_location("search_export", ROOT / "scripts" / "build-search-export-flow.py")
FLOW = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(FLOW)


def walk_actions(actions, depth=1, in_loop=False):
    for name, action in actions.items():
        yield name, action, actions, depth, in_loop
        nested_loop = in_loop or action["type"] in ("Foreach", "Until")
        yield from walk_actions(action.get("actions", {}), depth + 1, nested_loop)
        yield from walk_actions(action.get("else", {}).get("actions", {}), depth + 1, nested_loop)


class SearchExportFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.policy = FLOW.load_policy()
        cls.template = (FLOW.TARGET / "RuntimeSearchResults.template.xlsx").read_bytes()
        cls.definition = FLOW.build_definition(cls.policy, cls.template)
        cls.actions = {name: action for name, action, *_ in walk_actions(cls.definition["actions"])}

    def test_saved_definition_and_connector_contract(self):
        self.assertEqual(self.definition, json.loads((FLOW.TARGET / "definition.json").read_text()))
        trigger = self.definition["triggers"]["manual"]
        self.assertEqual((trigger["type"], trigger["kind"]), ("Request", "Skills"))
        self.assertEqual(set(trigger["inputs"]["schema"]["properties"]), {"query", "scope"})
        references = FLOW.connection_references()
        self.assertEqual(len(references), 5)
        self.assertTrue(all(value["runtimeSource"] == "invoker" for value in references.values()))
        self.assertFalse(self.definition["metadata"]["defaultToEmbeddedConnections"])
        responses = [a for a in self.actions.values() if a["type"] == "Response"]
        self.assertEqual(responses, [self.actions["Respond_to_agent"]])
        self.assertNotIn("Asynchronous", json.dumps(self.definition))
        self.assertEqual(self.actions["Respond_to_agent"]["runAfter"], {"Format_chat_result": ["Succeeded"]})
        self.assertEqual(self.actions["Continue_private_export"]["runAfter"], {"Respond_to_agent": ["Succeeded"]})

    def test_platform_structure_and_all_expression_references(self):
        entries = list(walk_actions(self.definition["actions"]))
        names = [name for name, *_ in entries]
        self.assertEqual(len(names), len(set(names)))
        self.assertLessEqual(max(depth for _, _, _, depth, _ in entries), 8)
        for name, action, siblings, _, in_loop in entries:
            with self.subTest(action=name):
                self.assertFalse(in_loop and action["type"] == "Terminate")
                self.assertTrue(set(action.get("runAfter", {})) <= set(siblings))
        text = json.dumps(self.definition)
        references = re.findall(r"(?:outputs|body|actions|items|result)\('([^']+)'\)", text)
        self.assertFalse(set(references) - set(names))
        variables = {v["name"] for a in self.actions.values() if a["type"] == "InitializeVariable"
                     for v in a["inputs"]["variables"]}
        self.assertFalse(set(re.findall(r"variables\('([^']+)'\)", text)) - variables)

    def test_150_separate_sites_and_multiple_sites_per_department(self):
        policy = copy.deepcopy(self.policy)
        policy["sites"] = [
            {**self.policy["sites"][0], "siteId": str(uuid.UUID(int=i + 1)),
             "url": f"{policy['tenantOrigin']}/sites/Offline-{i}",
             "department": ("HR", "Finance", "IT")[i // 50]}
            for i in range(150)
        ]
        policy["searchSiteUrl"] = policy["sites"][0]["url"]
        scopes = FLOW.compile_scopes(policy)
        self.assertEqual(len(scopes["All"]), 8)
        for scope, expected in (("All", 150), ("HR", 50), ("Finance", 50), ("IT", 50)):
            ids = [identifier for batch in scopes[scope] for identifier in batch["Sites"]]
            self.assertEqual(len(set(ids)), expected)
            for batch in scopes[scope]:
                self.assertLessEqual(len(batch["Sites"]), 20)
                self.assertIn("DepartmentId:" + policy["hubSiteCollectionId"], batch["Kql"])
                for identifier in batch["Sites"]:
                    self.assertIn(f'SiteID:"{identifier}"', batch["Kql"])
        policy["sites"].append(policy["sites"][0])
        with self.assertRaisesRegex(ValueError, "only once"):
            FLOW.compile_scopes(policy)
        policy["sites"] = [
            {**policy["sites"][0], "siteId": str(uuid.UUID(int=i + 1))}
            for i in range(241)
        ]
        with self.assertRaisesRegex(ValueError, "site-batch budget"):
            FLOW.compile_scopes(policy)

    def test_search_fetches_locators_for_files_and_pages_not_folders(self):
        request = FLOW.search_request("item()?['Kql']", 0)["request"]
        self.assertIsInstance(request["SelectProperties"], list)
        self.assertEqual(request["SelectProperties"], ["SPWebUrl", "SiteID", "ListID", "ListItemID"])
        self.assertEqual(request["RowLimit"], 100)
        self.assertIn("IsDocument:1 OR FileExtension:aspx", request["Querytext"])
        self.assertIn("STS_ListItem_WebPageLibrary", request["Querytext"])
        self.assertFalse(request["TrimDuplicates"])
        self.assertFalse(request["EnableQueryRules"])
        self.assertIn("body('Valid_current_items')", self.actions["Select_verified_rows"]["inputs"]["from"])
        self.assertIn("contains(body('Group_ids'),item()?['Id'])", self.actions["Valid_current_items"]["inputs"]["where"])
        self.assertEqual(self.actions["Read_current_items"]["inputs"]["parameters"]["dataset"],
                         "@items('Each_metadata_group')?['WebUrl']")
        self.assertIn("uriPath", self.actions["Scoped_locators"]["inputs"]["where"])
        self.assertIn("'/'", self.actions["Scoped_locators"]["inputs"]["where"])
        self.assertNotIn("FieldValuesAsText", json.dumps(self.definition))

    def test_preview_is_bounded_and_uses_current_permission_checked_metadata(self):
        self.assertEqual((FLOW.PREVIEW_ROWS, FLOW.PREVIEW_CANDIDATES), (5, 20))
        budget = self.actions["Preview_candidate_budget"]["expression"]
        self.assertIn("length(variables('PreviewRows')),5", budget)
        self.assertIn("variables('PreviewChecked'),20", budget)
        self.assertEqual(self.actions["Read_preview_item"]["runAfter"], {"Preview_field_names": ["Succeeded"]})
        self.assertEqual(self.actions["Preview_item_verified"]["runAfter"], {"Read_preview_item": ["Succeeded"]})
        self.assertIn("body('Read_preview_item')?['Id']", self.actions["Preview_item_verified"]["expression"])
        self.assertIn("ServerRelativeUrl", self.actions["Preview_item_verified"]["expression"])
        row = self.actions["Preview_row"]["inputs"]
        self.assertIn("body('Read_preview_item')?['TopicTags']", row["Tags"])
        self.assertIn("ServerRelativeUrl", row["URL"])
        self.assertEqual(self.actions["Remember_preview_line"]["runAfter"], {"Remember_preview_url": ["Succeeded"]})
        self.assertEqual(self.actions["Format_chat_result"]["runAfter"]["Build_chat_preview"], FLOW.ALL_STATES)

    def test_chat_has_two_columns_links_stored_tags_and_escaped_text(self):
        message = self.actions["Format_chat_result"]["inputs"]
        self.assertIn("| File or page | Stored tags |", message)
        self.assertIn("| --- | --- |", message)
        self.assertNotIn("| Type |", message)
        line = self.actions["Remember_preview_line"]["inputs"]["value"]
        self.assertIn("'](',outputs('Preview_row')?['URL'],') | '", line)
        self.assertIn("Not supplied", line)
        self.assertIn("full tags in Excel", line)
        self.assertNotIn("['Type']", line)
        escaped = FLOW.markdown_text("item()")
        for escape in ("%0D", "%0A", "%5C", "'|'", "'['", "']'", "&lt;", "&gt;"):
            self.assertIn(escape, escaped)

    def test_fresh_template_has_safe_row_specific_link_formula(self):
        scratch = ROOT / "tests" / ("workbook-check-" + uuid.uuid4().hex)
        try:
            with patch.object(FLOW, "TARGET", scratch):
                data = FLOW.create_runtime_template()
        finally:
            if scratch.exists():
                shutil.rmtree(scratch)
        book = load_workbook(io.BytesIO(data))
        try:
            sheet = book["Results"]
            table = sheet.tables["SearchResults"]
            self.assertEqual(table.ref, "A8:H9")
            self.assertEqual(sheet["A9"].value, FLOW.SENTINEL)
            self.assertEqual(sheet["G9"].value, FLOW.LINK_FORMULA)
            self.assertEqual(table.tableColumns[6].calculatedColumnFormula.attr_text, FLOW.LINK_FORMULA[1:])
            self.assertIn("SearchResults[[#This Row],[SourceURL]]", FLOW.LINK_FORMULA)
            self.assertNotIn("[@SourceURL]", FLOW.LINK_FORMULA)
            self.assertTrue(sheet.column_dimensions["H"].hidden)
            self.assertTrue(sheet.column_dimensions["F"].hidden)
            self.assertFalse(sheet["A9"].font.bold)
            self.assertEqual(sheet.freeze_panes, "A9")
        finally:
            book.close()
        for action in ("Replace_placeholder", "Append_result"):
            item = self.actions[action]["inputs"]["parameters"]["item"]
            self.assertEqual(item["URL"], FLOW.LINK_FORMULA)
            self.assertIn("['URL']", item["SourceURL"])
            self.assertIn("decodeUriComponent('%27')", item["SourceURL"])
        self.assertIn("equals(variables('Exported'),0)", self.actions["Write_result"]["expression"])

    def test_full_export_caps_verification_privacy_and_recipient(self):
        export = json.dumps(self.actions["Export_rows"])
        for limit in ("1000", "2000", "40"):
            self.assertIn(limit, export)
        verified = self.actions["Written_rows_verified"]["expression"]
        self.assertIn("intersection(body('Written_projection'),variables('ExpectedRows'))", verified)
        self.assertIn("union(body('Written_urls'),body('Written_urls'))", verified)
        self.assertEqual(self.actions["Delivery_acl_verified"]["runAfter"], {"Final_file_acl": ["Succeeded"]})
        self.assertIn("body('Personal_site_user')?['Id']", self.actions["Delivery_acl_verified"]["expression"])
        for action in self.actions.values():
            inputs = action.get("inputs", {})
            if isinstance(inputs, dict) and inputs.get("host", {}).get("operationId") == "SendEmailV2":
                self.assertEqual(inputs["parameters"]["emailMessage/To"], "@body('Selected_profile')?['mail']")
        self.assertIn("body('Source_directory_profile')?['id']", self.actions["Accounts_aligned"]["expression"])
        self.assertIn("body('Personal_drive')?['owner']?['user']?['id']", self.actions["Accounts_aligned"]["expression"])
        self.assertIn("Copy", str(self.actions["Make_report_private"]["inputs"]["parameters"]).replace("copy", "Copy"))
        self.assertIn("copyRoleAssignments=false", self.actions["Make_report_private"]["inputs"]["parameters"]["parameters/uri"])
        self.assertIn("'Partial'", json.dumps(self.actions["Complete_metadata_values"]))
        self.assertIn("1000", self.actions["Format_chat_result"]["inputs"])
        self.assertIn("not all matches", self.actions["Send_private_workbook"]["inputs"]["parameters"]["emailMessage/Body"])


if __name__ == "__main__":
    unittest.main()
