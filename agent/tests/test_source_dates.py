import io
import json
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from test_search_export_flow import FLOW, ROOT, walk_actions
from workbook_layout import HEADERS, HEADER_ROW, apply_layout, records


class SourceDateTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.definition = FLOW.build_definition(
            FLOW.load_policy(), (FLOW.TARGET / "RuntimeSearchResults.template.xlsx").read_bytes(),
        )
        cls.actions = {name: action for name, action, *_ in walk_actions(cls.definition["actions"])}

    def test_dates_come_only_from_verified_current_file_metadata(self):
        for action in ("Read_preview_item", "Read_current_items"):
            uri = self.actions[action]["inputs"]["parameters"]["parameters/uri"]
            self.assertIn("File/TimeLastModified,File/TimeCreated", uri)
            self.assertIn("$expand=File", uri)
            self.assertNotIn("LastModifiedTime", uri)
        preview = self.actions["Preview_row"]["inputs"]
        export = self.actions["Select_verified_rows"]["inputs"]["select"]
        for row, file in ((preview, "body('Read_preview_item')?['File']"), (export, "item()?['File']")):
            self.assertEqual(row["ModifiedUTC"], f"@coalesce({file}?['TimeLastModified'],'')")
            self.assertEqual(row["CreatedUTC"], f"@coalesce({file}?['TimeCreated'],'')")
        self.assertEqual(list(export), FLOW.HEADERS + list(FLOW.DATE_FIELDS))
        self.assertEqual(
            FLOW.search_request("item()?['Kql']", 0)["request"]["SelectProperties"],
            ["SPWebUrl", "SiteID", "ListID", "ListItemID"],
        )
        self.assertNotIn("PolicyStatus", json.dumps(export))
        self.assertNotIn("ReviewDate", json.dumps(export))

    def test_chat_dates_are_labelled_utc_null_safe_and_do_not_require_html(self):
        line = self.actions["Remember_preview_line"]["inputs"]["value"]
        message = self.actions["Format_chat_result"]["inputs"]
        self.assertIn("Created and last modified dates are UTC (YYYY-MM-DD).", message)
        self.assertIn("| File or page | Created (UTC) | Modified (UTC) | Stored tags |", message)
        self.assertNotIn("— Created: ", line)
        self.assertNotIn("; Modified: ", line)
        self.assertNotIn("<br", line)
        self.assertNotIn("\n", line)
        for field in ("CreatedUTC", "ModifiedUTC"):
            value = f"outputs('Preview_row')?['{field}']"
            self.assertEqual(
                FLOW.preview_date(field),
                f"if(empty({value}),'Not supplied',formatDateTime({value},'yyyy-MM-dd'))",
            )
            self.assertIn(FLOW.preview_date(field), line)
        with self.assertRaises(ValueError):
            FLOW.preview_date("GeneratedUTC")

    def test_full_timestamps_and_blank_semantics_survive_writer_and_readback(self):
        for action in ("Replace_placeholder", "Append_result"):
            item = self.actions[action]["inputs"]["parameters"]["item"]
            self.assertEqual(set(item), set(HEADERS) | set(FLOW.DATE_FIELDS) | {"SourceURL"})
            for field in ("CreatedUTC", "ModifiedUTC"):
                self.assertEqual(
                    item[field],
                    "@concat(decodeUriComponent('%27'),string(items('Each_verified_row')?['" + field + "']))",
                )
        written = self.actions["Written_projection"]["inputs"]["select"]
        for field in ("CreatedUTC", "ModifiedUTC"):
            self.assertEqual(written[field], "@coalesce(item()?['" + field + "'],'')")
        self.assertIn("ExpectedRows", self.actions["Written_rows_verified"]["expression"])

    def test_blank_templates_keep_existing_columns_and_row_relative_formula(self):
        for path, expected in (
            (ROOT / "exports" / "template" / "CorpNetSearchResults.template.xlsx", list(HEADERS)),
            (FLOW.TARGET / "RuntimeSearchResults.template.xlsx", list(HEADERS) + list(FLOW.DATE_FIELDS) + ["SourceURL"]),
        ):
            with self.subTest(template=path.name):
                book = load_workbook(path)
                try:
                    sheet = book["Results"]
                    table = sheet.tables["SearchResults"]
                    self.assertEqual([sheet.cell(HEADER_ROW, i).value for i in range(1, len(expected) + 1)], expected)
                    self.assertEqual([column.name for column in table.tableColumns], expected)
                    self.assertEqual(sheet.freeze_panes, "B9")
                    self.assertEqual(sheet["H9"].number_format, "@")
                    self.assertEqual(sheet["H9"].font.name, "Arial")
                    self.assertFalse(sheet["H9"].font.bold)
                    self.assertEqual(sheet.column_dimensions["H"].hidden, len(expected) > len(HEADERS))
                    self.assertIn("File.TimeCreated", sheet["H8"].comment.text)
                    if len(expected) > len(HEADERS):
                        self.assertEqual(sheet["G9"].value, FLOW.LINK_FORMULA)
                        self.assertEqual(table.tableColumns[6].calculatedColumnFormula.attr_text, FLOW.LINK_FORMULA[1:])
                        self.assertTrue(sheet.column_dimensions["K"].hidden)
                finally:
                    book.close()

    def make_book(self, columns=8):
        book = Workbook()
        sheet = book.active
        sheet.title = "Results"
        sheet.append(HEADERS[:columns])
        site = "https://contoso.sharepoint.com/sites/Fictional"
        examples = [
            ["Fictional file", "HR", "example;synthetic", "Document", "2026-09-12T23:59:59.1234567Z",
             site, site + "/Documents/example.docx", "2024-02-29T00:00:00Z"],
            ["Fictional page", "HR", "", "Page", "", site, site + "/SitePages/example.aspx", ""],
        ]
        for row in examples:
            sheet.append(row[:columns])
        sheet.add_table(Table(displayName="SearchResults", ref=f"A1:{chr(64 + columns)}3"))
        book.create_sheet("ExportInfo").append(["Field", "Value"])
        return book

    def test_full_precision_dates_and_missing_values_roundtrip_without_inference(self):
        book = self.make_book()
        expected = [values for _, values in records(book)]
        apply_layout(book)
        buffer = io.BytesIO()
        book.save(buffer)
        book.close()
        saved = load_workbook(io.BytesIO(buffer.getvalue()))
        try:
            self.assertEqual([values for _, values in records(saved)], expected)
            self.assertEqual(saved["Results"]["E9"].data_type, "s")
            self.assertEqual(saved["Results"]["H9"].data_type, "s")
            self.assertEqual(saved["Results"]["G9"].hyperlink.target, expected[0][6])
            self.assertIsNone(saved["Results"]["H10"].value)
        finally:
            saved.close()

    def test_historical_seven_column_layout_keeps_original_data(self):
        book = self.make_book(columns=7)
        try:
            expected = records(book)
            apply_layout(book)
            self.assertEqual([v for _, v in records(book)], [v for _, v in expected])
            self.assertEqual(book["Results"].tables["SearchResults"].ref, "A8:G10")
            self.assertEqual(book["Results"].freeze_panes, "B9")
        finally:
            book.close()

    def test_created_field_cannot_bypass_literal_source_validation(self):
        book = self.make_book()
        try:
            book["Results"]["H2"] = "=1+1"
            with self.assertRaisesRegex(ValueError, "literal, error-free"):
                records(book)
        finally:
            book.close()

    def test_compact_runtime_uses_true_excel_date_columns_and_hidden_lossless_sources(self):
        book = load_workbook(FLOW.TARGET / "RuntimeSearchResults.template.xlsx")
        try:
            sheet = book["Results"]
            self.assertEqual(sheet.tables["SearchResults"].ref, "A8:K9")
            for column in ("E", "F", "H", "K"):
                self.assertTrue(sheet.column_dimensions[column].hidden)
                self.assertFalse(sheet[f"{column}9"].alignment.wrap_text)
            for column, (label, field) in enumerate(FLOW.DATE_FIELDS.items(), 9):
                self.assertEqual(sheet.cell(8, column).value, label)
                self.assertEqual(sheet.cell(9, column).value, FLOW.excel_date_formula(field))
                self.assertEqual(sheet.cell(9, column).number_format, "yyyy-mm-dd")
                self.assertFalse(sheet.cell(9, column).alignment.wrap_text)
                self.assertEqual(sheet.tables["SearchResults"].tableColumns[column - 1].calculatedColumnFormula.attr_text,
                                 FLOW.excel_date_formula(field)[1:])
            self.assertEqual(book.epoch.isoformat(), "1899-12-30T00:00:00")
            self.assertEqual(book.calculation.calcMode, "auto")
        finally:
            book.close()

    def test_display_dates_are_written_as_formulas_and_verified_as_calendar_dates(self):
        projected = self.actions["Select_verified_rows"]["inputs"]["select"]
        written = self.actions["Written_projection"]["inputs"]["select"]
        for label, field in FLOW.DATE_FIELDS.items():
            for action in ("Replace_placeholder", "Append_result"):
                self.assertEqual(self.actions[action]["inputs"]["parameters"]["item"][label], FLOW.excel_date_formula(field))
            self.assertIn("formatDateTime", projected[label])
            self.assertIn("yyyy-MM-dd", projected[label])
            self.assertEqual(written[label], "@" + FLOW.written_date(label))
            self.assertIn("isFloat", written[label])
            self.assertIn("1899-12-30", written[label])
            self.assertIn("Not supplied", written[label])
        with self.assertRaises(ValueError):
            FLOW.excel_date_formula("GeneratedUTC")
        with self.assertRaises(ValueError):
            FLOW.written_date("GeneratedUTC")

    def test_masthead_is_compact_and_summary_references_actual_export_metadata(self):
        book = load_workbook(FLOW.TARGET / "RuntimeSearchResults.template.xlsx")
        try:
            sheet = book["Results"]
            self.assertLessEqual(sum(sheet.row_dimensions[i].height for i in range(1, 9)), 140)
            self.assertLessEqual(sheet.row_dimensions[9].height, 30)
            self.assertLessEqual(sheet["A1"].font.sz, 18)
            self.assertTrue(all(cell.alignment.vertical == "top" for cell in sheet[9]))
            for field in ("Scope", "Query"):
                self.assertIn(FLOW.metadata_formula(field), sheet["B4"].value)
            for field in ("ExportedRowCount", "FileCount", "PageCount", "CompletionStatus"):
                self.assertIn(FLOW.metadata_formula(field), sheet["B5"].value)
            self.assertNotIn("Full query and scope: ExportInfo", sheet["B4"].value)
        finally:
            book.close()

    def test_hidden_urls_and_raw_timestamps_do_not_inflate_local_row_height(self):
        book = self.make_book()
        try:
            sheet = book["Results"]
            sheet["G2"] = "https://contoso.sharepoint.com/sites/Fictional/" + "nested/" * 35 + "example.docx"
            before = [values for _, values in records(book)]
            apply_layout(book)
            self.assertLessEqual(sheet.row_dimensions[9].height, 30)
            self.assertEqual([values for _, values in records(book)], before)
            self.assertFalse(sheet["F9"].alignment.wrap_text)
            self.assertFalse(sheet["E9"].alignment.wrap_text)
            self.assertFalse(sheet["H9"].alignment.wrap_text)
        finally:
            book.close()

    def test_ten_preview_rows_keep_twenty_candidate_and_hundred_search_page_bounds(self):
        self.assertEqual((FLOW.PREVIEW_ROWS, FLOW.PREVIEW_CANDIDATES, FLOW.PAGE_SIZE), (10, 20, 100))
        for name in ("Preview_candidate_budget", "Preview_batch_budget"):
            expression = self.actions[name]["expression"]
            self.assertIn("length(variables('PreviewRows')),10", expression)
            self.assertIn("variables('PreviewChecked'),20", expression)
        self.assertIn("up to 10 results from at most 20 checked candidates", self.actions["Format_chat_result"]["inputs"])


if __name__ == "__main__":
    unittest.main()
