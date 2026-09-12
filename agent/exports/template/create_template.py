from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table
from workbook_layout import HEADERS, HEADER_ROW, apply_layout


OUTPUT = Path(__file__).with_name("CorpNetSearchResults.template.xlsx")


def create_template(path: Path) -> None:
    workbook = Workbook()
    workbook.properties.title = "CorpNet Search Results - Export Template"
    workbook.properties.creator = "CorpNet Search Hub"
    workbook.properties.description = "Blank private search-results export template."
    results = workbook.active
    results.title = "Results"
    results.append(HEADERS)
    results.append([None] * len(HEADERS))

    table = Table(
        displayName="SearchResults",
        ref=f"A1:{get_column_letter(len(HEADERS))}2",
        autoFilter=AutoFilter(ref=f"A1:{get_column_letter(len(HEADERS))}2"),
    )
    results.add_table(table)

    info = workbook.create_sheet("ExportInfo")
    rows = [
        ("CorpNet Search Results", "Private, permission-checked search export"),
        ("Query", None),
        ("Scope", None),
        ("RequestedBy", None),
        ("GeneratedUTC", None),
        ("CompletionStatus", "TEMPLATE - no search or export has run"),
        (
            "CountSemantics",
            "The delivery message must distinguish estimated index matches from rows "
            "actually exported and disclose any limit, omission or incomplete batch.",
        ),
        (
            "Sharing",
            "Create a separate workbook per request. Restrict access to the verified "
            "requesting user; do not use anonymous or organization-wide links.",
        ),
        (
            "Snapshot",
            "The report reflects export time. Original file links remain subject to "
            "current SharePoint permissions.",
        ),
        (
            "Tags",
            "Tags are stored source values. Use Excel Text Filters > Contains to "
            "filter a tag within a semicolon-delimited value.",
        ),
        (
            "TemplateInstructions",
            "Native AddRow writers must insert results before removing the blank placeholder "
            "so the new rows inherit body formatting, not the header. "
            "Update the table range and set real cell hyperlinks. Do not email this "
            "blank template as a completed report.",
        ),
        (
            "TextSafety",
            "Treat source titles, tags and other text as literal data, never executable "
            "formulas. Text formatting alone is not proof that a particular writer "
            "prevents formula injection.",
        ),
    ]
    for row in rows:
        info.append(row)
    apply_layout(workbook)
    workbook.save(path)


def validate_template(path: Path) -> None:
    workbook = load_workbook(path)
    results = workbook["Results"]
    assert tuple(cell.value for cell in results[HEADER_ROW]) == HEADERS
    assert results.freeze_panes == f"B{HEADER_ROW + 1}"
    assert results.tables["SearchResults"].autoFilter.ref == f"A{HEADER_ROW}:{get_column_letter(len(HEADERS))}{HEADER_ROW + 1}"
    assert all(cell.value is None for cell in results[HEADER_ROW + 1])
    assert workbook["ExportInfo"]["B6"].value.startswith("TEMPLATE")
    assert not any(cell.data_type in {"f", "e"} for sheet in workbook for row in sheet for cell in row)
    assert not any(cell.hyperlink for sheet in workbook for row in sheet for cell in row)
    workbook.close()


if __name__ == "__main__":
    create_template(OUTPUT)
    validate_template(OUTPUT)
    print(f"Created and verified blank filterable Excel template: {OUTPUT.name}")
