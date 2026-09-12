"""Shared presentation for the stable seven-column search export contract."""

import argparse
import math
import re
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import TableStyleInfo


HEADERS = ("Title", "Department", "Tags", "Type", "ModifiedUTC", "SourceSite", "URL")
HEADER_ROW = 8
LINK_FORMAT = ';;;"Open file"'
WIDTHS = (40, 16, 32, 14, 20, 38, 14)
NAVY = "17365D"
INK = "25364B"
BLUE = "0563C1"
PALE = "EAF1F8"


def table_bounds(book):
    return range_boundaries(book["Results"].tables["SearchResults"].ref)


def records(book):
    sheet = book["Results"]
    first_column, header, last_column, last = table_bounds(book)
    if (first_column, last_column) != (1, 7):
        raise ValueError("Expected the seven-column SearchResults table")
    if tuple(sheet.cell(header, column).value for column in range(1, 8)) != HEADERS:
        raise ValueError("The search export columns changed")
    rows = []
    for index in range(header + 1, last + 1):
        values = [sheet.cell(index, column).value or "" for column in range(1, 8)]
        if not any(values):
            continue
        url_cell = sheet.cell(index, 7)
        if url_cell.data_type == "f":
            match = re.fullmatch(r'=HYPERLINK\("((?:[^"]|"")*)","(?:[^"]|"")*"\)', values[6])
            if not match:
                raise ValueError("Only the existing literal-URL HYPERLINK formula can be converted")
            values[6] = match.group(1).replace('""', '"')
        site, target = urlsplit(values[5]), urlsplit(values[6])
        if (site.scheme != "https" or target.scheme != "https"
                or not site.hostname or not site.hostname.endswith(".sharepoint.com")
                or site.netloc.lower() != target.netloc.lower() or site.username or site.password
                or not target.path.lower().startswith(site.path.rstrip("/").lower() + "/")):
            raise ValueError("A source hyperlink is outside its SharePoint site")
        if any(sheet.cell(index, column).data_type in {"f", "e"} for column in range(1, 7)):
            raise ValueError("Source fields must remain literal, error-free values")
        rows.append((index, values))
    return rows


def _text(cell, value):
    cell.value = value
    cell.data_type = "s"


def apply_layout(book):
    sheet = book["Results"]
    _, old_header, _, old_last = table_bounds(book)
    if old_header not in (1, HEADER_ROW):
        raise ValueError("Unexpected table position; refusing to move unrelated worksheet content")
    records(book)
    if old_header == 1:
        sheet.insert_rows(1, HEADER_ROW - 1)
    last = old_last + HEADER_ROW - old_header
    table = sheet.tables["SearchResults"]
    table.ref = f"A{HEADER_ROW}:G{last}"
    table.autoFilter = AutoFilter(ref=table.ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    info = book["ExportInfo"]
    metadata = {row[0].value: row[1].value for row in info if row[0].value}
    mode = str(metadata.get("Mode", ""))
    badge = "OWNER-ONLY SAMPLE | NOT A FULL SEARCH" if mode.startswith("OWNER_ONLY") else "SEARCH EXPORT"
    for area in ("A1:G2", "A3:G3", "B4:G4", "B5:G5", "A6:G6"):
        if area not in sheet.merged_cells:
            sheet.merge_cells(area)
    _text(sheet["A1"], "CorpNet Search Results")
    _text(sheet["A3"], badge)
    for row, label, value in (
        (4, "Scope", metadata.get("Scope") or "Set by the search when the report is generated"),
        (5, "Status", metadata.get("CompletionStatus") or "See ExportInfo for completion details"),
    ):
        _text(sheet.cell(row, 1), label.upper())
        shown = str(value)
        if len(shown) > 180:
            shown = shown[:180] + "... (full details in ExportInfo)"
        _text(sheet.cell(row, 2), shown)
    _text(sheet["A6"], "Use the header filters to narrow results. Select Open file to view the original. Full export details: ExportInfo tab.")
    for row in sheet.iter_rows(min_row=1, max_row=HEADER_ROW, max_col=7):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
    for row in sheet.iter_rows(min_row=1, max_row=2, max_col=7):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].font = Font(name="Arial", size=23, bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(vertical="center", indent=1)
    for cell in sheet[3]:
        cell.fill = PatternFill("solid", fgColor=PALE)
    sheet["A3"].font = Font(name="Arial", size=10, bold=True, color=NAVY)
    sheet["A3"].alignment = Alignment(vertical="center", indent=2)
    for row in (4, 5):
        sheet.cell(row, 1).font = Font(name="Arial", size=9, bold=True, color="61758C")
        sheet.cell(row, 1).alignment = Alignment(vertical="center", indent=1)
    sheet["A6"].font = Font(name="Arial", size=10, color="61758C")
    sheet["A6"].alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    for row, height in {1: 18, 2: 34, 3: 26, 4: 34, 5: 30, 6: 32, 7: 9, 8: 28}.items():
        sheet.row_dimensions[row].height = height
    for column, width in zip("ABCDEFG", WIDTHS):
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["F"].hidden = True
    sheet.column_dimensions["G"].hidden = False
    sheet.column_dimensions["G"].number_format = LINK_FORMAT
    for cell in sheet[HEADER_ROW]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    sheet.cell(HEADER_ROW, 5).comment = Comment("Last modified time in UTC. Blank means the source did not supply it.", "CorpNet Search Hub")
    sheet.cell(HEADER_ROW, 7).comment = Comment("Open file links to the original. The full URL remains in the cell value; SourceSite is preserved in hidden column F.", "CorpNet Search Hub")
    bottom = Border(bottom=Side(style="hair", color="DCE5EF"))
    for index in range(HEADER_ROW + 1, last + 1):
        for column in range(1, 8):
            cell = sheet.cell(index, column)
            cell.font = Font(name="Arial", size=11, bold=False, color=INK)
            cell.fill = PatternFill()
            cell.border = bottom
            cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            cell.number_format = "@"
        sheet.cell(index, 7).number_format = LINK_FORMAT
        sheet.cell(index, 7).font = Font(name="Arial", size=11, color=BLUE, underline="single")
        sheet.cell(index, 7).alignment = Alignment(vertical="center", indent=1)
        lines = max(
            math.ceil(len(str(sheet.cell(index, column).value or "")) / (WIDTHS[column - 1] * .9))
            for column in range(1, 6)
        )
        sheet.row_dimensions[index].height = min(96, max(40, 14 * lines + 10))
    for index, values in records(book):
        link = sheet.cell(index, 7)
        _text(link, values[6])
        link.hyperlink = values[6]
        link.hyperlink.tooltip = "Open the original file; SharePoint checks your current access"
    sheet.freeze_panes = f"B{HEADER_ROW + 1}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_view.tabSelected = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = f"1:{HEADER_ROW}"
    sheet.print_area = f"A1:G{last}"
    sheet.oddFooter.center.text = "CorpNet Search Hub | Source permissions still apply"
    sheet.oddFooter.right.text = "Page &P of &N"
    for side in ("left", "right", "top", "bottom"):
        setattr(sheet.page_margins, side, .3)
    for row in info:
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.column == 1, color=INK)
            cell.fill = PatternFill("solid", fgColor="F3F6FA" if cell.column == 1 else "FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            cell.border = bottom
        info.row_dimensions[row[0].row].height = min(90, max(30, 15 * math.ceil(len(str(row[1].value or "")) / 85) + 12))
    for cell in info[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 100
    info.sheet_view.showGridLines = False
    info.sheet_view.zoomScale = 90
    info.sheet_view.tabSelected = False
    info.freeze_panes = "B2"
    book.active = book.sheetnames.index("Results")
    book.properties.title = "CorpNet Search Results"


def redesign(source, output):
    if output.exists() or output.resolve() == source.resolve():
        raise ValueError("Create a separate redesigned copy; never overwrite the delivered evidence")
    book = load_workbook(source)
    expected = [values for _, values in records(book)]
    metadata = [(cell.coordinate, cell.value, cell.data_type) for row in book["ExportInfo"] for cell in row]
    apply_layout(book)
    if any(cell.data_type in {"f", "e"} for sheet in book for row in sheet for cell in row):
        raise ValueError("The redesigned workbook must contain no executable formulas or errors")
    with output.open("xb") as target:
        book.save(target)
    book.close()
    checked = load_workbook(output)
    if [values for _, values in records(checked)] != expected:
        raise ValueError("The redesign changed source data")
    if [(cell.coordinate, cell.value, cell.data_type) for row in checked["ExportInfo"] for cell in row] != metadata:
        raise ValueError("The redesign changed export metadata")
    checked.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    redesign(arguments.source, arguments.output)
    print(f"Created redesigned copy: {arguments.output}")
